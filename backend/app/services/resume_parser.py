"""
Resume Parser Service.

Handles complete resume processing pipeline:
1. File validation (extension, size, header integrity, encryption)
2. Safe text extraction (PDF via pypdf, DOCX via python-docx, table parsing, sanitization)
3. Structured information parsing:
   - Contact information (name, email, phone, linkedin, github, website, location)
   - Categorized skills extraction (languages, frameworks, cloud/devops, databases, tools, soft skills)
   - Education extraction (degrees, fields of study, institutions, graduation dates, GPA)
   - Project extraction (titles, descriptions, tech stacks, links)
   - Work experience extraction (job titles, companies, dates, highlights)
   - Executive summary / objective
4. Relational database storage (Resume, canonical Skills, ResumeSkill junction)
"""

import io
import os
import re
from typing import Any

import docx
import pypdf
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import FileUploadException
from app.models.resume import Resume, ResumeSkill
from app.models.skill import Skill

# Supported file extensions
ALLOWED_EXTENSIONS = {".pdf", ".docx"}

# Magic bytes signatures for validation
PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"


# ============================================================================
# 1. FILE VALIDATION & INTEGRITY
# ============================================================================

def validate_resume_upload(
    filename: str | None,
    content: bytes,
    max_bytes: int = 10 * 1024 * 1024,
) -> str:
    """
    Validate uploaded resume file for extension, size, and header integrity.

    Raises FileUploadException on any validation failure.
    Returns normalized lowercase extension (e.g. '.pdf' or '.docx').
    """
    if not filename or not filename.strip():
        raise FileUploadException("No file selected or filename is empty.")

    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise FileUploadException(
            f"Unsupported file type '{ext}'. Allowed types: {', '.join(sorted(ALLOWED_EXTENSIONS))}."
        )

    # Check file size limits
    if len(content) == 0:
        raise FileUploadException("Uploaded file is empty (0 bytes).")

    if len(content) > max_bytes:
        max_mb = max_bytes // (1024 * 1024)
        raise FileUploadException(
            f"File too large ({len(content) / (1024 * 1024):.1f}MB). Maximum allowed size is {max_mb}MB."
        )

    # Validate header integrity and readability
    if ext == ".pdf":
        if not content.startswith(PDF_MAGIC) and PDF_MAGIC not in content[:1024]:
            raise FileUploadException("Invalid or corrupt PDF file: missing PDF header signature.")
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise FileUploadException(
                    "Encrypted or password-protected PDF files are not supported. Please remove the password and try again."
                )
            # Verify we can access pages
            _ = len(reader.pages)
        except FileUploadException:
            raise
        except Exception as e:
            raise FileUploadException(f"Invalid or corrupt PDF file: unable to parse document structure ({str(e)}).") from e

    elif ext == ".docx":
        if not content.startswith(ZIP_MAGIC) and b"PK" not in content[:128]:
            raise FileUploadException("Invalid or corrupt DOCX file: missing DOCX/ZIP signature.")
        try:
            doc = docx.Document(io.BytesIO(content))
            # Verify we can access document elements
            _ = len(doc.paragraphs)
        except Exception as e:
            raise FileUploadException(f"Invalid or corrupt DOCX file: unable to read document structure ({str(e)}).") from e

    return ext


# ============================================================================
# 2. TEXT EXTRACTION & SANITIZATION
# ============================================================================

def extract_resume_text(content: bytes, ext: str) -> str:
    """
    Safely extract plain text from PDF or DOCX file content.
    Strips null bytes and normalizes whitespace.
    """
    text_parts: list[str] = []

    if ext == ".pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            for _page_num, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        except Exception as e:
            raise FileUploadException(f"Failed to extract text from PDF: {str(e)}") from e

    elif ext == ".docx":
        try:
            doc = docx.Document(io.BytesIO(content))
            # 1. Paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    text_parts.append(p_text)

            # 2. Tables (resumes often use tables for contact, skills, education)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        # De-duplicate adjacent identical cells from merged columns
                        deduped: list[str] = []
                        for cell in row_cells:
                            if not deduped or cell != deduped[-1]:
                                deduped.append(cell)
                        text_parts.append(" | ".join(deduped))
        except Exception as e:
            raise FileUploadException(f"Failed to extract text from DOCX: {str(e)}") from e

    raw_text = "\n".join(text_parts)

    # Sanitize: strip NUL bytes (\x00) which crash PostgreSQL text columns
    sanitized = raw_text.replace("\x00", "")

    # Normalize newlines and excess whitespace
    sanitized = re.sub(r"\r\n", "\n", sanitized)
    sanitized = re.sub(r"\n{3,}", "\n\n", sanitized)
    sanitized = sanitized.strip()

    if not sanitized:
        sanitized = "No extractable text found in resume."

    return sanitized


# ============================================================================
# 3. STRUCTURED RESUME PARSING
# ============================================================================

# Curated skill taxonomy
SKILL_TAXONOMY = {
    "languages": [
        "python", "javascript", "typescript", "java", "c++", "c#", "go", "golang",
        "rust", "ruby", "php", "swift", "kotlin", "sql", "html", "html5", "css",
        "css3", "bash", "shell", "r", "scala", "dart", "perl"
    ],
    "frameworks": [
        "react", "react.js", "next.js", "vue", "vue.js", "angular", "node.js",
        "express", "fastapi", "django", "flask", "spring boot", "spring",
        "asp.net", "ruby on rails", "rails", "pytorch", "tensorflow", "keras",
        "pandas", "numpy", "scikit-learn", "tailwind", "tailwind css", "bootstrap",
        "redux", "graphql", "rest", "restful"
    ],
    "cloud_devops": [
        "aws", "amazon web services", "azure", "gcp", "google cloud", "docker",
        "kubernetes", "terraform", "ansible", "ci/cd", "github actions", "gitlab ci",
        "jenkins", "linux", "nginx", "apache", "helm", "cloudformation", "prometheus",
        "grafana"
    ],
    "databases": [
        "postgresql", "postgres", "mysql", "mongodb", "redis", "elasticsearch",
        "sqlite", "dynamodb", "cassandra", "snowflake", "oracle", "mariadb",
        "firebase", "neo4j", "supabase"
    ],
    "tools": [
        "git", "github", "gitlab", "jira", "confluence", "postman", "figma",
        "vs code", "docker compose", "webpack", "vite", "npm", "yarn", "pnpm"
    ],
    "methodologies": [
        "microservices", "system design", "agile", "scrum", "tdd",
        "test driven development", "machine learning", "deep learning",
        "natural language processing", "nlp", "computer vision",
        "data engineering", "distributed systems", "oop", "clean architecture"
    ],
    "soft_skills": [
        "leadership", "communication", "problem solving", "teamwork",
        "collaboration", "mentorship", "project management", "critical thinking",
        "time management", "adaptability", "analytical thinking"
    ]
}


def extract_contact_info(text: str) -> dict[str, Any]:
    """Extract candidate contact details: name, email, phone, links, location."""
    contact: dict[str, Any] = {
        "name": None,
        "email": None,
        "phone": None,
        "linkedin": None,
        "github": None,
        "website": None,
        "location": None,
    }

    # 1. Email extraction
    email_match = re.search(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b", text)
    if email_match:
        contact["email"] = email_match.group(0).lower()

    # 2. Phone extraction (international and domestic formats)
    phone_match = re.search(
        r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b", text
    )
    if phone_match:
        contact["phone"] = phone_match.group(0).strip()

    # 3. LinkedIn
    linkedin_match = re.search(
        r"(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/([a-zA-Z0-9_\-\.]+)", text, re.IGNORECASE
    )
    if linkedin_match:
        contact["linkedin"] = f"https://linkedin.com/in/{linkedin_match.group(1)}"

    # 4. GitHub
    github_match = re.search(
        r"(?:https?:\/\/)?(?:www\.)?github\.com\/([a-zA-Z0-9_\-\.]+)", text, re.IGNORECASE
    )
    if github_match:
        contact["github"] = f"https://github.com/{github_match.group(1)}"

    # 5. Website / Portfolio
    website_match = re.search(
        r"(?:https?:\/\/)?(?:www\.)?([a-zA-Z0-9-]+\.(?:dev|io|me|portfolio|tech|com))\b",
        text,
        re.IGNORECASE,
    )
    if website_match:
        raw_site = website_match.group(0)
        if "linkedin.com" not in raw_site and "github.com" not in raw_site:
            contact["website"] = raw_site if raw_site.startswith("http") else f"https://{raw_site}"

    # 6. Candidate Name Extraction (heuristic from top lines)
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    skip_headers = {
        "resume", "curriculum vitae", "cv", "profile", "summary",
        "contact", "education", "experience", "skills", "projects"
    }
    for line in lines[:5]:
        cleaned_line = re.sub(r"[^\w\s]", "", line).strip()
        lower_line = line.lower()
        if (
            "@" not in line
            and "http" not in lower_line
            and "phone" not in lower_line
            and lower_line not in skip_headers
            and 2 <= len(cleaned_line.split()) <= 4
            and not any(char.isdigit() for char in line)
        ):
            contact["name"] = line.strip()
            break

    # 7. Location extraction
    location_match = re.search(
        r"\b([A-Z][a-zA-Z\s]+,\s*(?:[A-Z]{2}|[A-Z][a-zA-Z]+))\b", text
    )
    if location_match:
        loc = location_match.group(1).strip()
        # Avoid matching degrees as locations
        if not re.search(r"\b(Bachelor|Master|University|College|Science|Arts)\b", loc, re.I):
            contact["location"] = loc

    return contact


def extract_skills(text: str) -> dict[str, Any]:
    """Extract skills categorized by domain and return consolidated list."""
    text_lower = text.lower()
    found_categorized: dict[str, list[str]] = {cat: [] for cat in SKILL_TAXONOMY}
    all_found: list[str] = []

    for category, skill_list in SKILL_TAXONOMY.items():
        for skill in skill_list:
            # Escape skill for regex, handling special chars like c++, c#, .js, ci/cd
            escaped = re.escape(skill)
            # Use word boundaries or specific symbol boundaries
            pattern = rf"(?<![\w\.\+#]){escaped}(?![\w\.\+#])"
            if re.search(pattern, text_lower):
                # Capitalize nicely for display
                display_name = _format_skill_name(skill)
                if display_name not in found_categorized[category]:
                    found_categorized[category].append(display_name)
                if display_name not in all_found:
                    all_found.append(display_name)

    return {
        "all": all_found,
        "categories": found_categorized,
        "total_count": len(all_found),
    }


def _format_skill_name(raw: str) -> str:
    """Format canonical skill names with correct casing."""
    casing_map = {
        "python": "Python", "javascript": "JavaScript", "typescript": "TypeScript",
        "java": "Java", "c++": "C++", "c#": "C#", "go": "Go", "golang": "Go",
        "rust": "Rust", "ruby": "Ruby", "php": "PHP", "swift": "Swift",
        "kotlin": "Kotlin", "sql": "SQL", "html": "HTML", "html5": "HTML5",
        "css": "CSS", "css3": "CSS3", "bash": "Bash", "shell": "Shell",
        "react": "React", "react.js": "React", "next.js": "Next.js", "vue": "Vue",
        "vue.js": "Vue", "angular": "Angular", "node.js": "Node.js",
        "express": "Express", "fastapi": "FastAPI", "django": "Django",
        "flask": "Flask", "spring boot": "Spring Boot", "asp.net": "ASP.NET",
        "ruby on rails": "Ruby on Rails", "pytorch": "PyTorch", "tensorflow": "TensorFlow",
        "pandas": "Pandas", "numpy": "NumPy", "scikit-learn": "Scikit-Learn",
        "tailwind": "Tailwind CSS", "tailwind css": "Tailwind CSS",
        "redux": "Redux", "graphql": "GraphQL", "rest": "REST APIs",
        "aws": "AWS", "azure": "Azure", "gcp": "GCP", "docker": "Docker",
        "kubernetes": "Kubernetes", "terraform": "Terraform", "ansible": "Ansible",
        "ci/cd": "CI/CD", "github actions": "GitHub Actions", "jenkins": "Jenkins",
        "linux": "Linux", "nginx": "Nginx", "postgresql": "PostgreSQL",
        "postgres": "PostgreSQL", "mysql": "MySQL", "mongodb": "MongoDB",
        "redis": "Redis", "elasticsearch": "Elasticsearch", "sqlite": "SQLite",
        "dynamodb": "DynamoDB", "snowflake": "Snowflake", "git": "Git",
        "github": "GitHub", "gitlab": "GitLab", "jira": "Jira", "postman": "Postman",
        "figma": "Figma", "microservices": "Microservices", "agile": "Agile",
        "scrum": "Scrum", "tdd": "TDD", "leadership": "Leadership",
        "communication": "Communication", "problem solving": "Problem Solving",
        "teamwork": "Teamwork"
    }
    return casing_map.get(raw.lower(), raw.title())


def extract_education(text: str) -> list[dict[str, Any]]:
    """Extract education details: degree, field of study, institution, year, GPA."""
    education_entries: list[dict[str, Any]] = []

    # Look for education section
    edu_section = ""
    edu_match = re.search(
        r"(?:EDUCATION|ACADEMIC BACKGROUND|ACADEMIC QUALIFICATIONS)(.*?)(?:EXPERIENCE|PROJECTS|SKILLS|CERTIFICATIONS|PUBLICATIONS|$)",
        text,
        re.DOTALL | re.IGNORECASE,
    )
    edu_section = edu_match.group(1).strip() if edu_match else text

    # Regex patterns for degrees
    degree_patterns = [
        (r"\b(Ph\.?D\.?|Doctorate|Doctor of Philosophy)\b", "Doctor of Philosophy (Ph.D.)"),
        (r"\b(Master of Science|M\.?S\.?|Master of Arts|M\.?A\.?|MBA|Master's)\b", "Master of Science"),
        (r"\b(Bachelor of Science|B\.?S\.?|Bachelor of Arts|B\.?A\.?|B\.?Tech|B\.?E\.?|Bachelor's)\b", "Bachelor of Science"),
        (r"\b(Associate of Science|Associate Degree|A\.?S\.?)\b", "Associate Degree"),
    ]

    # Major/Field patterns
    major_match = re.search(
        r"\b(?:in|of)\s+([A-Z][a-zA-Z\s]+(?:Computer Science|Software Engineering|Data Science|Information Technology|Electrical Engineering|Computer Engineering|Mathematics|Physics|Business))\b",
        edu_section,
        re.IGNORECASE,
    )
    detected_major = major_match.group(1).strip() if major_match else None

    # Institution extraction: inspect lines containing educational institution keywords
    detected_institution = None
    for line in edu_section.split("\n"):
        line_clean = line.strip()
        if re.search(r"\b(University|College|Institute|Polytechnic|Academy|School)\b", line_clean, re.IGNORECASE):
            inst_candidate = re.split(r"\||–|\(|\d{4}", line_clean)[0].strip()
            inst_candidate = re.sub(
                r"^(?:Bachelor|Master|Doctor|Ph\.?D|Associate|B\.?S|M\.?S)[^,]*,\s*",
                "",
                inst_candidate,
                flags=re.IGNORECASE,
            ).strip()
            if inst_candidate:
                detected_institution = inst_candidate
                break

    # Year / Date pattern
    year_match = re.search(r"\b(20\d{2}\s*(?:-|–|to)\s*(?:20\d{2}|Present)|(?:May|June|July|August|Dec|December)?\s*20\d{2})\b", edu_section)
    detected_year = year_match.group(0).strip() if year_match else None

    # GPA pattern
    gpa_match = re.search(r"\bGPA:?\s*([0-4]\.\d{1,2}(?:\s*\/\s*4\.0)?)\b", edu_section, re.IGNORECASE)
    detected_gpa = gpa_match.group(1).strip() if gpa_match else None

    # Identify degrees found
    for pattern, normalized_degree in degree_patterns:
        if re.search(pattern, edu_section, re.IGNORECASE):
            education_entries.append({
                "degree": normalized_degree,
                "field": detected_major or "Computer Science",
                "institution": detected_institution or "Accredited University",
                "graduation_year": detected_year,
                "gpa": detected_gpa,
            })
            break

    # If no degree regex matched but university found, create entry
    if not education_entries and detected_institution:
        education_entries.append({
            "degree": "Higher Education Degree",
            "field": detected_major or "General Studies",
            "institution": detected_institution,
            "graduation_year": detected_year,
            "gpa": detected_gpa,
        })

    return education_entries


def extract_projects(text: str) -> list[dict[str, Any]]:
    """Extract projects: titles, descriptions, technologies, links."""
    projects: list[dict[str, Any]] = []

    # Locate projects section
    proj_match = re.search(
        r"(?:PROJECTS|PERSONAL PROJECTS|KEY PROJECTS|NOTABLE PROJECTS)(.*?)(?:EDUCATION|EXPERIENCE|SKILLS|CERTIFICATIONS|PUBLICATIONS|$)",
        text,
        re.DOTALL | re.IGNORECASE,
    )
    proj_section = proj_match.group(1).strip() if proj_match else ""

    if proj_section:
        # Split project section by lines or bullet points
        blocks = re.split(r"\n(?=[A-Z0-9][A-Za-z0-9\s\-_]{3,40}(?:\s*\||\s*–|\s*\(|\s*\n))", proj_section)
        for block in blocks:
            lines = [line.strip() for line in block.split("\n") if line.strip()]
            if not lines:
                continue

            title_line = lines[0]
            # Strip dates, URLs, or pipes from title
            clean_title = re.split(r"\||–|-|\(|http", title_line)[0].strip()

            if len(clean_title) >= 3 and not re.search(r"^(github|technologies|tools):", clean_title, re.I):
                # Extract link if present
                link_match = re.search(r"(https?:\/\/[^\s\)]+)", block)
                link = link_match.group(1) if link_match else None

                # Extract technologies mentioned in block
                block_skills = extract_skills(block)["all"]

                description = " ".join(lines[1:]) if len(lines) > 1 else lines[0]

                projects.append({
                    "title": clean_title,
                    "description": description[:300],
                    "technologies": block_skills[:8],
                    "link": link,
                })

    return projects


def extract_experience(text: str) -> list[dict[str, Any]]:
    """Extract professional experience: title, company, dates, highlights."""
    experiences: list[dict[str, Any]] = []

    exp_match = re.search(
        r"(?:WORK EXPERIENCE|PROFESSIONAL EXPERIENCE|EXPERIENCE|EMPLOYMENT HISTORY)(.*?)(?:EDUCATION|PROJECTS|SKILLS|CERTIFICATIONS|PUBLICATIONS|$)",
        text,
        re.DOTALL | re.IGNORECASE,
    )
    exp_section = exp_match.group(1).strip() if exp_match else ""

    if exp_section:
        # Common job titles
        title_patterns = [
            r"((?:Senior\s+|Lead\s+|Principal\s+|Junior\s+|Staff\s+)?(?:Software Engineer|Full Stack Developer|Frontend Developer|Backend Developer|DevOps Engineer|Data Scientist|Machine Learning Engineer|Cloud Architect|Product Manager|System Administrator|Engineering Manager))"
        ]

        blocks = re.split(r"\n(?=[A-Z][A-Za-z0-9\s,]{3,50}(?:\s*\||\s*–|\s*-\s*|\s*20\d{2}))", exp_section)
        for block in blocks:
            lines = [line.strip() for line in block.split("\n") if line.strip()]
            if not lines:
                continue

            header_line = lines[0]
            detected_title = "Software Engineer"
            for tp in title_patterns:
                m = re.search(tp, header_line, re.IGNORECASE)
                if m:
                    detected_title = m.group(1).title()
                    break

            # Company name extraction
            company_match = re.search(r"(?:at|@|\||,)\s*([A-Z][A-Za-z0-9\s]+?)(?:\s*\||\s*–|\s*-|\s*\(|\s*\d{4}|$)", header_line)
            company = company_match.group(1).strip() if company_match else "Technology Company"

            # Dates
            date_match = re.search(r"((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)?\s*\d{4}\s*(?:-|–|to)\s*(?:Present|\w+\s*\d{4}|\d{4}))", block, re.IGNORECASE)
            dates = date_match.group(1).strip() if date_match else "Recent"

            # Highlights / bullets
            highlights = [line.lstrip("•-* ").strip() for line in lines[1:] if len(line.strip()) > 15]

            experiences.append({
                "title": detected_title,
                "company": company,
                "dates": dates,
                "highlights": highlights[:4],
            })

    return experiences


def extract_summary(text: str) -> str | None:
    """Extract professional summary or objective statement."""
    summary_match = re.search(
        r"(?:SUMMARY|PROFESSIONAL SUMMARY|EXECUTIVE SUMMARY|OBJECTIVE|ABOUT ME)(.*?)(?:EXPERIENCE|EDUCATION|SKILLS|PROJECTS|$)",
        text,
        re.DOTALL | re.IGNORECASE,
    )
    if summary_match:
        summary_text = summary_match.group(1).strip()
        lines = [line.strip() for line in summary_text.split("\n") if line.strip()]
        if lines:
            return " ".join(lines[:3])[:500]
    return None


# ============================================================================
# 4. ORCHESTRATION & DATABASE STORAGE
# ============================================================================

def parse_full_resume(raw_text: str) -> dict[str, Any]:
    """
    Orchestrate full resume extraction and parsing.
    Returns contact_info and comprehensive parsed_data structure.
    """
    contact_info = extract_contact_info(raw_text)
    skills_data = extract_skills(raw_text)
    education_data = extract_education(raw_text)
    projects_data = extract_projects(raw_text)
    experience_data = extract_experience(raw_text)
    summary_data = extract_summary(raw_text)

    text_lower = raw_text.lower()
    sections_detected = {
        "contact": bool(contact_info.get("email") or contact_info.get("phone")),
        "summary": bool(summary_data),
        "skills": skills_data["total_count"] > 0,
        "experience": "experience" in text_lower or len(experience_data) > 0,
        "education": "education" in text_lower or len(education_data) > 0,
        "projects": "project" in text_lower or len(projects_data) > 0,
    }

    parsed_data = {
        "summary": summary_data,
        "skills": skills_data,
        "education": education_data,
        "projects": projects_data,
        "experience": experience_data,
        "sections_detected": sections_detected,
        "word_count": len(raw_text.split()),
    }

    return {
        "contact_info": contact_info,
        "parsed_data": parsed_data,
    }


async def save_parsed_resume_to_db(
    db: AsyncSession,
    resume: Resume,
    raw_text: str,
    parsed_result: dict[str, Any],
) -> Resume:
    """
    Persist parsed resume details and synchronize extracted skills to DB.
    """
    resume.raw_text = raw_text
    resume.contact_info = parsed_result["contact_info"]
    resume.parsed_data = parsed_result["parsed_data"]
    resume.status = "parsed"
    db.add(resume)
    await db.flush()

    # Link canonical skills and ResumeSkill junction records
    extracted_skills: list[str] = parsed_result["parsed_data"]["skills"]["all"]
    categories: dict[str, list[str]] = parsed_result["parsed_data"]["skills"]["categories"]

    # Inverted lookup for category
    skill_to_cat: dict[str, str] = {}
    for cat, sk_list in categories.items():
        for sk in sk_list:
            skill_to_cat[sk.lower()] = cat

    for skill_name in extracted_skills[:40]:  # limit to top 40 skills per resume
        cat = skill_to_cat.get(skill_name.lower(), "technical")

        # Check if Skill exists
        existing_skill_result = await db.execute(
            select(Skill).where(Skill.name.ilike(skill_name))
        )
        canonical_skill = existing_skill_result.scalar_one_or_none()
        if not canonical_skill:
            canonical_skill = Skill(
                name=skill_name,
                category=cat,
                is_active=True,
            )
            db.add(canonical_skill)
            await db.flush()
            await db.refresh(canonical_skill)

        # Check if ResumeSkill already linked
        existing_link_result = await db.execute(
            select(ResumeSkill).where(
                ResumeSkill.resume_id == resume.id,
                ResumeSkill.skill_id == canonical_skill.id,
            )
        )
        if not existing_link_result.scalar_one_or_none():
            resume_skill = ResumeSkill(
                resume_id=resume.id,
                skill_id=canonical_skill.id,
                proficiency="intermediate",
            )
            db.add(resume_skill)

    await db.flush()
    await db.refresh(resume)
    return resume
