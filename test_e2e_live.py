import urllib.request
import json
import io
import sys

BASE = "http://localhost:8000/api/v1"

def req(endpoint, method="GET", data=None, token=None, files=None):
    url = f"{BASE}{endpoint}"
    headers = {}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    
    body = None
    if files:
        boundary = "----WebKitFormBoundary7MA4YWxkTrZu0gW"
        headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
        buf = io.BytesIO()
        for field, (fname, fcontent, ftype) in files.items():
            buf.write(f"--{boundary}\r\n".encode())
            buf.write(f'Content-Disposition: form-data; name="{field}"; filename="{fname}"\r\n'.encode())
            buf.write(f"Content-Type: {ftype}\r\n\r\n".encode())
            buf.write(fcontent if isinstance(fcontent, bytes) else fcontent.encode("utf-8"))
            buf.write(b"\r\n")
        buf.write(f"--{boundary}--\r\n".encode())
        body = buf.getvalue()
    elif data is not None:
        headers["Content-Type"] = "application/json"
        body = json.dumps(data).encode("utf-8")
    
    request = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        raw = e.read().decode("utf-8")
        try:
            return e.code, json.loads(raw)
        except Exception:
            return e.code, raw

def main():
    print("=" * 60)
    print("INTEGRATION VERIFICATION: ALL 10 ENDPOINTS")
    print("=" * 60)

    # 1. Health Probe
    s, r = req("/health")
    assert s == 200, f"Health probe failed: {s}"
    print(f"1. [PASS] Health Probe: status={r.get('status')}")

    # 2. User Registration
    import time
    email = f"alex.verify.{int(time.time())}@career.com"
    s, r = req("/auth/register", "POST", {"email": email, "password": "Password123!", "full_name": "Alex Verify"})
    assert s == 201, f"Registration failed: {s} - {r}"
    print(f"2. [PASS] User Registration: created {r['email']}")

    # 3. User Login
    s, r = req("/auth/login", "POST", {"email": email, "password": "Password123!"})
    assert s == 200, f"Login failed: {s} - {r}"
    token = r["access_token"]
    print(f"3. [PASS] User Login: JWT token acquired ({r['token_type']})")

    # 4. Authenticated Profile
    s, r = req("/auth/me", token=token)
    assert s == 200, f"Profile fetch failed: {s}"
    print(f"4. [PASS] Profile: {r['full_name']} ({r['email']})")

    # 5. List Resumes
    s, r = req("/resumes", token=token)
    assert s == 200, f"List resumes failed: {s}"
    print(f"5. [PASS] List Resumes: count={len(r)}")

    # 6. Upload Resume (valid docx)
    import docx
    doc = docx.Document()
    doc.add_heading("Alex Developer", 0)
    doc.add_paragraph("Senior Software Engineer with 5+ years of experience in distributed systems.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, FastAPI, Docker, PostgreSQL, React, Next.js, LangChain, Redis")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Senior Backend Engineer at TechCorp. Built scalable microservices and optimized PostgreSQL database queries.")
    doc_buf = io.BytesIO()
    doc.save(doc_buf)
    valid_docx = doc_buf.getvalue()

    s, r = req("/resumes/upload", "POST", token=token, files={"file": ("engineer_resume.docx", valid_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert s == 201, f"Upload resume failed: {s} - {r}"
    resume_id = r["id"]
    print(f"6. [PASS] Upload Resume: id={resume_id}, filename={r['original_filename']}")

    # 7. Analyze Resume ATS
    s, r = req(f"/resumes/{resume_id}/analyze", "POST", {}, token=token)
    assert s == 200, f"ATS Analysis failed: {s} - {r}"
    print(f"7. [PASS] ATS Analysis: overall_score={r['overall_score']}")

    # 8. Job Recommendations
    s, r = req("/jobs/recommendations", token=token)
    assert s == 200, f"Job recommendations failed: {s}"
    print(f"8. [PASS] Job Recommendations: retrieved {len(r)} matching roles")

    # 9. Create Chat Session & Send Message
    s, r = req("/chat/sessions", "POST", {"title": "Full Stack Advisory"}, token=token)
    assert s == 201, f"Create chat session failed: {s}"
    session_id = r["id"]
    print(f"9. [PASS] Chat Session: created session {session_id}")

    s, r = req(f"/chat/sessions/{session_id}/messages", "POST", {"content": "Hello, what skills should I learn next?"}, token=token)
    assert s in (200, 201), f"Send message failed: {s} - {r}"
    safe_content = r['content'][:65].encode("ascii", "replace").decode("ascii")
    print(f"10. [PASS] Advisor Response: {safe_content}...")

    print("=" * 60)
    print("ALL VERIFICATIONS PASSED: 10/10 SUCCESS")
    print("=" * 60)

if __name__ == "__main__":
    main()
